"""
Document ingestor — scans a folder and indexes files into the shared RAG.

Supported: PDF, DOCX, XLSX, TXT, MD, CSV, RTF
Each file gets chunked and stored in ChromaDB with source_type="document".
"""

import hashlib
import logging
import multiprocessing
import os
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Iterator

if TYPE_CHECKING:
    from services.rag_engine import RAGEngine

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv", ".rtf"}
_EXTRACT_TIMEOUT = 30  # seconds per file


@dataclass
class DocIngestProgress:
    status: str = "idle"   # idle | running | completed | error
    processed: int = 0
    total: int = 0
    message: str = ""


_progress = DocIngestProgress()


def get_progress() -> DocIngestProgress:
    return _progress


def _doc_id(path: Path) -> str:
    h = hashlib.sha1(str(path).encode()).hexdigest()[:16]
    return f"doc:{h}"


def _is_mostly_garbage(text: str) -> bool:
    """True when pdfminer returned almost no real prose — likely a scanned PDF."""
    import re
    stripped = text.strip()
    if not stripped or len(stripped) < 50:
        return True
    # Strip UUIDs, page-break chars, and known watermark patterns before counting words
    cleaned = re.sub(r'[A-F0-9]{8}-[A-F0-9]{4}-[A-F0-9]{4}-[A-F0-9]{4}-[A-F0-9]{12}', '', stripped, flags=re.IGNORECASE)
    cleaned = re.sub(r'Authentisign\s+ID\s*:', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'[\x0c\x00-\x08\x0b\x0e-\x1f]', ' ', cleaned)  # control chars / form feeds
    words = re.findall(r'[A-Za-z]{4,}', cleaned)
    return len(words) < 8


_OCR_MAX_PAGES = 15   # enough for most agreements; limits pdftoppm time
_OCR_DPI       = 100  # 100 DPI → 4× faster than 200 DPI; sufficient for text


def _find_tesseract() -> str | None:
    """Locate the tesseract executable across platforms. shutil.which covers
    the common case (on PATH); these are just the well-known install
    locations for each OS when it isn't."""
    import shutil
    found = shutil.which("tesseract")
    if found:
        return found
    candidates = (
        [r"C:\Program Files\Tesseract-OCR\tesseract.exe",
         r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"]
        if sys.platform == "win32"
        else ["/opt/homebrew/bin/tesseract", "/usr/local/bin/tesseract", "/usr/bin/tesseract"]
    )
    for c in candidates:
        if os.path.exists(c):
            return c
    return None


def _ocr_pdf(path_str: str) -> str:
    """OCR up to _OCR_MAX_PAGES pages of a scanned PDF. Returns combined text."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
        if pytesseract.pytesseract.tesseract_cmd == "tesseract":
            found = _find_tesseract()
            if found:
                pytesseract.pytesseract.tesseract_cmd = found
            else:
                logger.warning(f"[docs] tesseract not found - cannot OCR {path_str}")
                return ""
        images = convert_from_path(path_str, dpi=_OCR_DPI, last_page=_OCR_MAX_PAGES)
        pages = []
        for img in images:
            page_text = pytesseract.image_to_string(img, lang="eng")
            if page_text.strip():
                pages.append(page_text.strip())
        return "\n\n".join(pages)
    except Exception as e:
        logger.warning(f"[docs] OCR failed for {path_str}: {e}")
        return ""


def _extract_worker_loop(req_queue: "multiprocessing.Queue", resp_queue: "multiprocessing.Queue") -> None:
    """Runs in a persistent child process, one file at a time from the queue -
    completely isolated from the server, but reused across the whole ingest
    run instead of respawned per file (spawning re-imports the whole app,
    including the embedding model, in the child on Windows since there's no
    fork() there - that made every single file pay a multi-second penalty).
    Calls os.setsid() so a hung file's SIGKILL can also reap pdftoppm
    grandchildren spawned by pdf2image during OCR.
    """
    import os as _os
    try:
        _os.setsid()
    except Exception:
        pass
    while True:
        path_str = req_queue.get()
        if path_str is None:  # sentinel: shut down
            return
        try:
            path = Path(path_str)
            ext = path.suffix.lower()
            text = ""
            if ext == ".pdf":
                from pdfminer.high_level import extract_text
                text = extract_text(path_str) or ""
                if _is_mostly_garbage(text):
                    text = _ocr_pdf(path_str)
            elif ext == ".docx":
                import docx
                doc = docx.Document(path_str)
                text = "\n".join(p.text for p in doc.paragraphs)
            elif ext == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(path_str, read_only=True, data_only=True)
                parts = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        line = "\t".join(str(c) if c is not None else "" for c in row)
                        if line.strip():
                            parts.append(line)
                text = "\n".join(parts)
            elif ext in (".txt", ".md", ".csv", ".rtf"):
                text = path.read_text(errors="replace")
            resp_queue.put(("ok", text))
        except Exception as e:
            resp_queue.put(("error", str(e)))


_PLAIN_TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".rtf"}


class _ExtractWorkerPool:
    """Owns the single persistent extraction worker process. Restarts it
    lazily only when it's never been started or died/hung on the previous
    file - not on every call."""

    def __init__(self):
        self._proc: "multiprocessing.Process | None" = None
        self._req_q = None
        self._resp_q = None

    def _ensure_started(self):
        if self._proc is not None and self._proc.is_alive():
            return
        ctx = multiprocessing.get_context("spawn")
        self._req_q = ctx.Queue()
        self._resp_q = ctx.Queue()
        self._proc = ctx.Process(
            target=_extract_worker_loop, args=(self._req_q, self._resp_q), daemon=True
        )
        self._proc.start()

    def _kill(self):
        import os as _os
        import signal as _signal
        if self._proc is not None:
            try:
                _os.killpg(_os.getpgid(self._proc.pid), _signal.SIGKILL)
            except Exception:
                try:
                    self._proc.kill()
                except Exception:
                    pass
            self._proc.join(timeout=3)
        self._proc = None
        self._req_q = None
        self._resp_q = None

    def extract(self, path: Path) -> str:
        self._ensure_started()
        self._req_q.put(str(path))
        try:
            status, payload = self._resp_q.get(timeout=_EXTRACT_TIMEOUT)
        except Exception:
            logger.warning(f"[docs] extract timed out after {_EXTRACT_TIMEOUT}s: {path.name}")
            self._kill()
            return ""
        if status == "error":
            logger.warning(f"[docs] extract failed {path.name}: {payload}")
            return ""
        return payload


_worker_pool = _ExtractWorkerPool()


def _extract_text(path: Path) -> str:
    """Extract text for one file. Plain-text formats are read inline (can't
    hang the way a PDF/DOCX/XLSX parser might); everything else goes through
    the persistent extraction worker process so a hung file can be hard-killed
    without taking the whole ingest run down."""
    if path.suffix.lower() in _PLAIN_TEXT_EXTENSIONS:
        try:
            return path.read_text(errors="replace")
        except OSError as e:
            logger.warning(f"[docs] extract failed {path.name}: {e}")
            return ""
    return _worker_pool.extract(path)


def _iter_files(folder: Path) -> Iterator[Path]:
    try:
        for p in folder.rglob("*"):
            try:
                if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS:
                    yield p
            except (PermissionError, OSError):
                continue
    except (PermissionError, OSError) as e:
        logger.warning(f"[docs] cannot read {folder}: {e}")


def ingest_folder(folder_path: str, rag: "RAGEngine") -> int:
    """Scan folder lazily and ingest new/modified files. Returns newly indexed count."""
    global _progress
    folder = Path(folder_path).expanduser().resolve()
    if not folder.exists():
        _progress = DocIngestProgress(status="error", message=f"Folder not found: {folder}")
        return 0

    _progress = DocIngestProgress(status="running", message=f"Scanning {folder.name}…")

    scanned = 0
    new_count = 0

    for path in _iter_files(folder):
        scanned += 1
        _progress.processed = scanned
        _progress.message = f"Checking {path.name}…"

        try:
            mtime = str(int(path.stat().st_mtime))
        except OSError:
            continue

        doc_id = _doc_id(path)

        if rag.is_document_current(doc_id, mtime):
            continue

        _progress.message = f"Indexing {path.name}…"
        try:
            text = _extract_text(path)
        except Exception as e:
            logger.warning(f"[docs] skipping {path.name}: {e}")
            continue
        if not text.strip():
            continue

        ok = rag.ingest_document(
            doc_id=doc_id,
            text=text,
            filename=path.name,
            file_path=str(path),
            file_type=path.suffix.lstrip(".").lower(),
            modified_at=mtime,
        )
        if ok:
            new_count += 1
        else:
            logger.warning(f"[docs] storage upsert failed for {path.name} - not counted as indexed")

    _progress = DocIngestProgress(
        status="completed",
        processed=scanned,
        total=scanned,
        message=f"Done — {new_count} new file{'s' if new_count != 1 else ''} indexed ({scanned} scanned)",
    )
    return new_count
